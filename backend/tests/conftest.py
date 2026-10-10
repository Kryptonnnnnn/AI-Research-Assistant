"""Test harness: isolated SQLite DB + upload dir per test, hash embeddings, mock LLM. No network, no paid APIs."""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

import httpx
import pytest
import pytest_asyncio

_TMP = tempfile.mkdtemp(prefix="ara-tests-")
os.environ.update(
    APP_ENV="test",
    DATABASE_URL=f"sqlite+aiosqlite:///{_TMP}/base.db",
    STORAGE_LOCAL_DIR=f"{_TMP}/uploads",
    EMBEDDING_PROVIDER="hash",
    LLM_PROVIDER="mock",
    JOB_RUNNER="disabled",
    RATE_LIMIT_ENABLED="false",
    SECRET_KEY="test-secret-key-test-secret-key-test-secret-key",
    JOB_RETRY_BACKOFF_SECONDS="0",
    RETRIEVAL_SCORE_THRESHOLD="0.1",
)

import app.models  # noqa: E402,F401
from app.core.config import reset_settings_cache  # noqa: E402
from app.core.ratelimit import limiter  # noqa: E402
from app.db.base import Base, dispose_engine, get_engine, get_sessionmaker  # noqa: E402
from app.embeddings.factory import set_embedding_provider  # noqa: E402
from app.jobs import queue  # noqa: E402
from app.llm.factory import set_llm_override  # noqa: E402
from app.storage.factory import set_storage  # noqa: E402

PASSWORD = "Str0ngPassw0rd!"


PG_URL = os.environ.get("TEST_DATABASE_URL")  # e.g. postgresql+asyncpg://ara:...@localhost:5432/ara_test


def _prepare_postgres_once() -> None:
    """Run the real Alembic migrations once per session against the test Postgres (also tests the migrations)."""
    import subprocess
    import sys

    env = {**os.environ, "DATABASE_URL": PG_URL}
    root = Path(__file__).resolve().parent.parent
    subprocess.run([sys.executable, "-m", "alembic", "downgrade", "base"], cwd=root, env=env, check=True)
    subprocess.run([sys.executable, "-m", "alembic", "upgrade", "head"], cwd=root, env=env, check=True)


@pytest.fixture(scope="session", autouse=True)
def _pg_session():
    if PG_URL:
        _prepare_postgres_once()


@pytest_asyncio.fixture
async def app(tmp_path):
    os.environ["DATABASE_URL"] = PG_URL or f"sqlite+aiosqlite:///{tmp_path}/test.db"
    os.environ["STORAGE_LOCAL_DIR"] = str(tmp_path / "uploads")
    reset_settings_cache()
    set_storage(None)
    set_embedding_provider(None)
    set_llm_override(None)
    limiter.reset()
    await dispose_engine()
    async with get_engine().begin() as conn:
        if PG_URL:
            tables = ", ".join(t.name for t in Base.metadata.sorted_tables)
            await conn.execute(__import__("sqlalchemy").text(f"TRUNCATE {tables} CASCADE"))
        else:
            await conn.run_sync(Base.metadata.create_all)
    from app.main import create_app

    yield create_app()
    set_llm_override(None)
    await dispose_engine()


def make_client(app) -> httpx.AsyncClient:
    return httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")


@pytest_asyncio.fixture
async def client(app):
    async with make_client(app) as c:
        yield c


async def register(c: httpx.AsyncClient, email: str = "alice@example.com", password: str = PASSWORD) -> dict:
    r = await c.post("/api/v1/auth/register", json={"email": email, "password": password, "display_name": email.split("@")[0]})
    assert r.status_code == 201, r.text
    return r.json()


@pytest_asyncio.fixture
async def alice(client):
    await register(client, "alice@example.com")
    return client


@pytest_asyncio.fixture
async def bob(app):
    async with make_client(app) as c:
        await register(c, "bob@example.com")
        yield c


async def process_all() -> int:
    return await queue.drain(get_sessionmaker())


async def upload(c, name: str, data: bytes, mime: str = "text/plain"):
    return await c.post("/api/v1/documents/upload", files=[("files", (name, data, mime))])


async def upload_and_process(c, name: str, data: bytes, mime: str = "text/plain") -> dict:
    r = await upload(c, name, data, mime)
    assert r.status_code == 201, r.text
    item = r.json()["items"][0]
    assert item["status"] == "created", item
    await process_all()
    doc = (await c.get(f"/api/v1/documents/{item['document']['id']}")).json()
    return doc


# ---- sample file builders ----
def make_pdf(pages: list[str]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=A4)
    for text in pages:
        y = 800
        for line in text.split("\n"):
            cv.drawString(60, y, line)
            y -= 16
        cv.showPage()
    cv.save()
    return buf.getvalue()


def make_blank_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    cv = canvas.Canvas(buf, pagesize=A4)
    cv.rect(50, 50, 200, 200, fill=1)  # image-like content, no text
    cv.showPage()
    cv.save()
    return buf.getvalue()


def make_docx(paragraphs: list[tuple[str, str]]) -> bytes:
    from docx import Document

    d = Document()
    for style, text in paragraphs:
        if style.startswith("h"):
            d.add_heading(text, level=int(style[1]))
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


PHOTO = (
    "Photosynthesis converts light energy into chemical energy stored in glucose. Chlorophyll in the chloroplasts "
    "absorbs red and blue light most strongly. The light-dependent reactions produce ATP and NADPH, which power the Calvin cycle."
)
TRANSFORMER = (
    "The Transformer architecture relies entirely on self-attention mechanisms instead of recurrence. Multi-head attention "
    "lets the model attend to information from different representation subspaces. Training was parallelised across eight GPUs."
)
FIXTURES = Path(__file__).parent
