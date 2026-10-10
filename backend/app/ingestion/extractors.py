"""Text extraction for PDF, DOCX, TXT and Markdown. Pure functions on bytes (CPU-bound)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field

from app.ingestion.cleaning import clean_text, unwrap_lines


class ExtractionError(Exception):
    """Permanent failure (do not retry). Message is safe to show to the user."""


@dataclass
class Segment:
    text: str
    page: int | None = None
    heading: str | None = None


@dataclass
class ExtractedDocument:
    title: str | None
    segments: list[Segment]
    page_count: int | None = None
    meta: dict = field(default_factory=dict)

    @property
    def total_chars(self) -> int:
        return sum(len(s.text) for s in self.segments)


MIN_TEXT_CHARS = 20


def extract(data: bytes, ext: str) -> ExtractedDocument:
    if ext == ".pdf":
        doc = extract_pdf(data)
    elif ext == ".docx":
        doc = extract_docx(data)
    elif ext in (".txt", ".md", ".markdown"):
        doc = extract_text(data, markdown=ext != ".txt")
    else:
        raise ExtractionError(f"Unsupported file type '{ext}'.")
    if doc.total_chars < MIN_TEXT_CHARS:
        if ext == ".pdf":
            raise ExtractionError(
                "No extractable text found. This looks like a scanned/image-only PDF; OCR is not supported "
                "in this app, so please upload a text-based PDF."
            )
        raise ExtractionError("The document contains no readable text.")
    return doc


def extract_pdf(data: bytes) -> ExtractedDocument:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data), strict=False)
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ExtractionError("The PDF is password-protected and cannot be read.")
            except ExtractionError:
                raise
            except Exception as exc:
                raise ExtractionError("The PDF is encrypted and cannot be read.") from exc
        pages = reader.pages
        n = len(pages)
        segments: list[Segment] = []
        for i, page in enumerate(pages, start=1):
            try:
                raw = page.extract_text() or ""
            except Exception:  # one bad page should not sink the whole document
                raw = ""
            text = unwrap_lines(clean_text(raw))
            if text:
                segments.append(Segment(text=text, page=i))
        title = None
        try:
            md = reader.metadata
            if md and md.title and str(md.title).strip():
                title = str(md.title).strip()[:500]
        except Exception:
            pass
        return ExtractedDocument(title=title, segments=segments, page_count=n)
    except ExtractionError:
        raise
    except (PdfReadError, ValueError, KeyError, OSError, AttributeError, TypeError) as exc:
        raise ExtractionError("The PDF appears to be corrupted or malformed.") from exc


def extract_docx(data: bytes) -> ExtractedDocument:
    from docx import Document as DocxDocument

    try:
        doc = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError("The DOCX file appears to be corrupted or malformed.") from exc

    segments: list[Segment] = []
    heading: str | None = None
    title: str | None = None
    buf: list[str] = []

    def flush() -> None:
        if buf:
            segments.append(Segment(text=clean_text("\n\n".join(buf)), heading=heading))
            buf.clear()

    for block in _iter_docx_blocks(doc):
        kind, text = block
        text = text.strip()
        if not text:
            continue
        if kind == "heading":
            flush()
            heading = text[:500]
            if title is None:
                title = heading
            buf.append(text)
        else:
            buf.append(text)
    flush()
    try:
        core_title = (doc.core_properties.title or "").strip()
        if core_title:
            title = core_title[:500]
    except Exception:
        pass
    return ExtractedDocument(title=title, segments=segments, page_count=None)


def _iter_docx_blocks(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            style = (p.style.name or "") if p.style is not None else ""
            if style.lower().startswith("heading") or style.lower() == "title":
                yield "heading", p.text
            else:
                yield "para", p.text
        elif tag == "tbl":
            table = Table(child, doc)
            rows = []
            for row in table.rows:
                cells: list[str] = []
                for c in row.cells:
                    t = c.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != t:  # merged cells repeat
                        cells.append(t)
                rows.append(" | ".join(cells))
            yield "para", "\n".join(rows)


def _decode(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):  # only trust UTF-16 when a BOM says so
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def extract_text(data: bytes, markdown: bool = False) -> ExtractedDocument:
    text = clean_text(_decode(data))
    if not markdown:
        first = next((ln.strip() for ln in text.split("\n") if ln.strip()), None)
        return ExtractedDocument(
            title=(first[:120] if first and len(first) < 120 else None), segments=[Segment(text=text)] if text else []
        )
    segments: list[Segment] = []
    heading: str | None = None
    title: str | None = None
    buf: list[str] = []
    in_code = False

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            segments.append(Segment(text=body, heading=heading))
        buf.clear()

    for line in text.split("\n"):
        if line.strip().startswith("```"):
            in_code = not in_code
        m = None if in_code else _MD_HEADING.match(line)
        if m:
            flush()
            heading = m.group(2).strip()[:500]
            if title is None and len(m.group(1)) == 1:
                title = heading
        buf.append(line)
    flush()
    return ExtractedDocument(title=title, segments=segments)
